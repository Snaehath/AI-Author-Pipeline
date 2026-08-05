"""
Manuscript Exporter Utility Module.

Converts generated novel manuscripts (outputs/<book_slug>/generated_novel_manuscript.md)
into beautifully styled DOCX (.docx) and PDF (.pdf) documents for human reading & evaluation.
"""

import json
import re
import sys
from pathlib import Path
from typing import Dict, Any, Optional
import sys

# Ensure project root is in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from AI_Author.utils.logger import setup_logger

logger = setup_logger("AI_Author.Utils.ManuscriptExporter")


def export_to_docx(markdown_path: Path, output_docx_path: Path) -> Path:
    """Converts Markdown manuscript file into a formatted DOCX document."""
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.warning("'python-docx' is not installed. Run: pip install python-docx")
        # Fallback raw write
        with open(markdown_path, "r", encoding="utf-8") as f:
            content = f.read()
        output_docx_path.write_text(content, encoding="utf-8")
        return output_docx_path

    from AI_Author.utils.text_sanitizer import sanitize_generated_prose

    doc = docx.Document()

    # Set Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    with open(markdown_path, "r", encoding="utf-8") as f:
        raw_content = f.read()

    sanitized_content = sanitize_generated_prose(raw_content)
    lines = sanitized_content.splitlines()

    for line in lines:
        text = line.strip()
        if not text:
            continue

        if text.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text.replace("# ", ""))
            run.font.name = "Georgia"
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        elif text.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(text.replace("## ", ""))
            run.font.name = "Georgia"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        elif text.startswith("*") and text.endswith("*"):
            p = doc.add_paragraph()
            run = p.add_run(text.strip("*"))
            run.font.name = "Georgia"
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        else:
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = "Georgia"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    doc.save(str(output_docx_path))
    logger.info(f"Exported DOCX manuscript to: '{output_docx_path}'")
    return output_docx_path


def sanitize_ascii(text: str) -> str:
    """Replaces Unicode em-dashes and smart quotes with standard ASCII equivalents."""
    replacements = {
        "—": "-",
        "–": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "…": "...",
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)
    return text.encode("latin-1", "ignore").decode("latin-1")


def export_to_pdf(markdown_path: Path, output_pdf_path: Path) -> Path:
    """Converts Markdown manuscript file into a formatted PDF document."""
    try:
        from fpdf import FPDF
    except ImportError:
        logger.warning("'fpdf2' is not installed. Run: pip install fpdf2")
        with open(markdown_path, "r", encoding="utf-8") as f:
            content = f.read()
        output_pdf_path.write_text(content, encoding="utf-8")
        return output_pdf_path

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(128, 128, 128)
            header_text = sanitize_ascii("AI Author Studio - Manuscript Evaluation Copy")
            self.cell(0, 10, header_text, border=0, align="R")
            self.ln(10)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, f"Page {self.page_no()}", border=0, align="C")

    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    from AI_Author.utils.text_sanitizer import sanitize_generated_prose

    with open(markdown_path, "r", encoding="utf-8") as f:
        raw_content = f.read()

    sanitized_content = sanitize_generated_prose(raw_content)
    lines = sanitized_content.splitlines()

    for line in lines:
        text = sanitize_ascii(line.strip())
        if not text:
            pdf.ln(4)
            continue

        if text.startswith("# "):
            pdf.set_font("Helvetica", "B", 18)
            pdf.set_text_color(17, 24, 39)
            pdf.cell(0, 12, text.replace("# ", ""), border=0, align="C")
            pdf.ln(6)
        elif text.startswith("## "):
            pdf.set_font("Helvetica", "B", 13)
            pdf.set_text_color(55, 65, 81)
            pdf.cell(0, 10, text.replace("## ", ""), border=0, align="L")
            pdf.ln(4)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(31, 41, 55)
            pdf.multi_cell(0, 6, text)
            pdf.ln(2)

    pdf.output(str(output_pdf_path))
    logger.info(f"Exported PDF manuscript to: '{output_pdf_path}'")
    return output_pdf_path


def export_book_workspace(book_dir: Path) -> Dict[str, str]:
    """Exports a book workspace to both .docx and .pdf files.

    Args:
        book_dir: Path to book output directory (e.g. outputs/generated_novel).

    Returns:
        Dictionary containing docx and pdf file paths.
    """
    b_path = Path(book_dir).resolve()
    md_file = b_path / "generated_novel_manuscript.md"
    if not md_file.exists():
        md_file = b_path / "cleaned.txt"

    if not md_file.exists():
        raise FileNotFoundError(f"No manuscript markdown or cleaned text file found in: '{b_path}'")

    docx_path = b_path / f"{b_path.name}_manuscript.docx"
    pdf_path = b_path / f"{b_path.name}_manuscript.pdf"

    export_to_docx(md_file, docx_path)
    export_to_pdf(md_file, pdf_path)

    return {
        "docx_path": str(docx_path),
        "pdf_path": str(pdf_path),
    }


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="AI Author Studio — Manuscript Exporter")
    parser.add_argument("book_dir", nargs="?", default="outputs/generated_novel", help="Path to book directory (default: outputs/generated_novel)")

    args = parser.parse_args()
    res = export_book_workspace(Path(args.book_dir))
    print(f"Export Complete!\n- DOCX: {res['docx_path']}\n- PDF: {res['pdf_path']}")
